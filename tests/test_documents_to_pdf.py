from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document
from openpyxl import Workbook
from pypdf import PdfReader

from lou_codi_pdf_converter.services.documents_to_pdf import convert_document_to_pdf


def test_office_conversion_invokes_libreoffice_headlessly(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    target = tmp_path / "result.pdf"
    source.touch()
    commands: list[list[str]] = []

    def fake_run(command: list[str], **_kwargs: object) -> Mock:
        commands.append(command)
        output_directory = Path(command[command.index("--outdir") + 1])
        (output_directory / "source.pdf").write_bytes(b"pdf")
        return Mock(returncode=0, stderr="", stdout="")

    with patch(
        "lou_codi_pdf_converter.services.documents_to_pdf.find_office_engine",
        return_value=Path("/opt/lou-codi/office-runtime/program/soffice"),
    ), patch(
        "lou_codi_pdf_converter.services.documents_to_pdf.subprocess.run",
        side_effect=fake_run,
    ):
        convert_document_to_pdf(source, target)

    assert target.read_bytes() == b"pdf"
    assert commands[0][0] == "/opt/lou-codi/office-runtime/program/soffice"
    assert "--headless" in commands[0]
    assert any(argument.startswith("-env:UserInstallation=file://") for argument in commands[0])


def test_docx_uses_lightweight_renderer_without_libreoffice(tmp_path: Path) -> None:
    source = tmp_path / "fiche.docx"
    target = tmp_path / "fiche.pdf"
    document = Document()
    document.add_heading("Fiche technique", level=1)
    document.add_paragraph("Une conversion locale légère.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Pièce"
    table.cell(0, 1).text = "Quantité"
    table.cell(1, 0).text = "Équerre"
    table.cell(1, 1).text = "4"
    document.save(source)

    with patch(
        "lou_codi_pdf_converter.services.documents_to_pdf.find_office_engine",
        return_value=None,
    ):
        engine = convert_document_to_pdf(source, target)

    assert engine == "lightweight"
    assert len(PdfReader(target).pages) == 1


def test_xlsx_uses_lightweight_renderer_without_libreoffice(tmp_path: Path) -> None:
    source = tmp_path / "mesures.xlsx"
    target = tmp_path / "mesures.pdf"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Mesures"
    sheet.append(["Repère", "Longueur"])
    sheet.append(["A", 125])
    workbook.save(source)

    with patch(
        "lou_codi_pdf_converter.services.documents_to_pdf.find_office_engine",
        return_value=None,
    ):
        engine = convert_document_to_pdf(source, target)

    assert engine == "lightweight"
    assert len(PdfReader(target).pages) == 1
