from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX, WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from lou_codi_pdf_converter.services.docx_import import import_docx_as_canvas


def test_docx_import_keeps_text_lists_and_character_styles(tmp_path: Path) -> None:
    source = tmp_path / "source.docx"
    document = Document()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Titre important")
    title_run.bold = True
    title_run.italic = True
    title_run.underline = True
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
    title_run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    bullet = document.add_paragraph(style="List Bullet")
    bullet.add_run("Premier point")
    document.save(source)

    imported = import_docx_as_canvas(source)

    root = imported.objects[0]
    assert root.text == "Titre important\n• Premier point"
    assert root.text_styles[0] == {
        "start": 0,
        "end": 15,
        "bold": True,
        "italic": True,
        "underline": True,
        "strikethrough": False,
        "font_family": "Sans-serif",
        "font_size": 16.0,
        "color": "#CC0000",
        "highlight": "#FFF176",
    }
    assert root.paragraph_alignments == [
        {"start": 0, "end": 15, "align": "center"}
    ]
    assert imported.page_count >= 1
