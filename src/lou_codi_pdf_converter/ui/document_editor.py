from __future__ import annotations

import json
import tempfile
import tkinter as tk
from html import escape
from pathlib import Path
from tkinter import filedialog, messagebox, simpledialog, ttk

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..i18n import translate_widget_tree
from ..paths import downloads_directory
from ..system import open_path
from .windowing import make_modal, maximize_window

NAVY = "#193B66"
BACKGROUND = "#F0F4F9"


class DocumentEditorWindow(tk.Toplevel):
    """A deliberately small local word processor centred on an A4 page."""

    def __init__(self, parent: tk.Misc) -> None:
        super().__init__(parent, class_="LouCodiPdfConverter")
        self.title("Créer un document — PDF Converter Lou Còdi")
        self.configure(bg=BACKGROUND)
        self.project_path: Path | None = None
        self.last_export: Path | None = None
        icon = getattr(self._root(), "app_icon", None)
        if icon is not None:
            self.iconphoto(False, icon)
        self._build()
        translate_widget_tree(self)
        maximize_window(self)
        make_modal(self, parent)

    def _build(self) -> None:
        header = tk.Frame(self, bg=NAVY, padx=22, pady=12)
        header.pack(fill="x")
        tk.Label(
            header,
            text="Créer un document",
            bg=NAVY,
            fg="white",
            font=("Sans", 19, "bold"),
        ).pack(side="left")
        tk.Label(
            header,
            text="Rédaction locale — page A4",
            bg=NAVY,
            fg="#CBDFF5",
            font=("Sans", 10),
        ).pack(side="left", padx=18)

        toolbar = ttk.Frame(self, style="App.TFrame", padding=(16, 10))
        toolbar.pack(fill="x")
        for label, command in (
            ("Nouveau", self._new),
            ("Ouvrir", self._load),
            ("Enregistrer", self._save),
            ("Gras", lambda: self._toggle_tag("bold")),
            ("Italique", lambda: self._toggle_tag("italic")),
            ("Souligné", lambda: self._toggle_tag("underline")),
            ("Tableau", self._insert_table),
            ("Exporter DOCX", self._export_docx),
            ("Exporter PDF", self._export_pdf),
            ("Imprimer", self._print),
        ):
            ttk.Button(toolbar, text=label, command=command).pack(side="left", padx=3)

        workspace = tk.Canvas(self, bg="#C8D0DA", highlightthickness=0)
        workspace.pack(fill="both", expand=True)
        scrollbar = ttk.Scrollbar(workspace, orient="vertical", command=workspace.yview)
        scrollbar.pack(side="right", fill="y")
        workspace.configure(yscrollcommand=scrollbar.set)

        page = tk.Frame(workspace, bg="white", width=794, height=1123, padx=64, pady=65)
        page.pack_propagate(False)
        page_window = workspace.create_window(0, 24, window=page, anchor="n")

        def centre_page(event: tk.Event[tk.Misc]) -> None:
            workspace.itemconfigure(page_window, x=max(event.width // 2, 420))
            workspace.configure(scrollregion=workspace.bbox("all"))

        workspace.bind("<Configure>", centre_page)
        self.editor = tk.Text(
            page,
            wrap="word",
            undo=True,
            maxundo=-1,
            font=("Liberation Serif", 12),
            relief="flat",
            bg="white",
            padx=2,
            pady=2,
        )
        self.editor.pack(fill="both", expand=True)
        self.editor.tag_configure("bold", font=("Liberation Serif", 12, "bold"))
        self.editor.tag_configure("italic", font=("Liberation Serif", 12, "italic"))
        self.editor.tag_configure("underline", underline=True)
        self.editor.focus_set()

        self.status = ttk.Label(self, text="Nouveau document", background=BACKGROUND)
        self.status.pack(fill="x", padx=18, pady=(4, 8))

    def _toggle_tag(self, tag: str) -> None:
        try:
            start, end = self.editor.index("sel.first"), self.editor.index("sel.last")
        except tk.TclError:
            return
        if tag in self.editor.tag_names(start):
            self.editor.tag_remove(tag, start, end)
        else:
            self.editor.tag_add(tag, start, end)

    def _new(self) -> None:
        self.editor.delete("1.0", "end")
        self.project_path = None
        self.status.configure(text="Nouveau document")

    def _insert_table(self) -> None:
        rows = simpledialog.askinteger("Insérer un tableau", "Nombre de lignes :", parent=self, minvalue=1, maxvalue=30)
        if rows is None:
            return
        columns = simpledialog.askinteger("Insérer un tableau", "Nombre de colonnes :", parent=self, minvalue=1, maxvalue=10)
        if columns is None:
            return
        line = "\t".join("Cellule" for _ in range(columns))
        self.editor.insert("insert", "\n" + "\n".join(line for _ in range(rows)) + "\n")

    def _project_data(self) -> dict[str, object]:
        tags = {
            tag: [(str(start), str(end)) for start, end in zip(self.editor.tag_ranges(tag)[::2], self.editor.tag_ranges(tag)[1::2])]
            for tag in ("bold", "italic", "underline")
        }
        return {"format": 1, "text": self.editor.get("1.0", "end-1c"), "tags": tags}

    def _save(self) -> None:
        path = self.project_path
        if path is None:
            selected = filedialog.asksaveasfilename(
                parent=self,
                title="Enregistrer le document Lou Còdi",
                initialdir=downloads_directory(),
                defaultextension=".loucodi-document",
                filetypes=(("Document Lou Còdi", "*.loucodi-document"),),
            )
            if not selected:
                return
            path = Path(selected)
        path.write_text(json.dumps(self._project_data(), ensure_ascii=False, indent=2), encoding="utf-8")
        self.project_path = path
        self.status.configure(text=f"Enregistré : {path.name}")

    def _load(self) -> None:
        selected = filedialog.askopenfilename(
            parent=self,
            title="Ouvrir un document Lou Còdi",
            initialdir=downloads_directory(),
            filetypes=(("Document Lou Còdi", "*.loucodi-document"),),
        )
        if not selected:
            return
        path = Path(selected)
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
            self.editor.delete("1.0", "end")
            self.editor.insert("1.0", str(data["text"]))
            for tag, ranges in data.get("tags", {}).items():
                if tag in {"bold", "italic", "underline"}:
                    for start, end in ranges:
                        self.editor.tag_add(tag, start, end)
        except (OSError, ValueError, KeyError, TypeError) as exc:
            messagebox.showerror("Document illisible", str(exc), parent=self)
            return
        self.project_path = path
        self.status.configure(text=f"Ouvert : {path.name}")

    def _choose_export(self, extension: str, label: str) -> Path | None:
        stem = self.project_path.stem if self.project_path else "nouveau_document"
        selected = filedialog.asksaveasfilename(
            parent=self,
            title=label,
            initialdir=downloads_directory(),
            initialfile=f"{stem}{extension}",
            defaultextension=extension,
            filetypes=((label, f"*{extension}"),),
        )
        return Path(selected) if selected else None

    def _lines(self) -> list[str]:
        return self.editor.get("1.0", "end-1c").splitlines()

    def _line_segments(self, line_number: int) -> list[tuple[str, frozenset[str]]]:
        active: set[str] = set()
        segments: list[tuple[str, frozenset[str]]] = []
        for kind, value, _index in self.editor.dump(
            f"{line_number}.0", f"{line_number}.end", text=True, tag=True
        ):
            if kind == "tagon" and value in {"bold", "italic", "underline"}:
                active.add(value)
            elif kind == "tagoff":
                active.discard(value)
            elif kind == "text" and value:
                segments.append((value, frozenset(active)))
        return segments

    def _export_docx(self) -> None:
        target = self._choose_export(".docx", "Document Word")
        if target is None:
            return
        document = Document()
        lines = self._lines()
        line_number = 1
        while line_number <= len(lines):
            line = lines[line_number - 1]
            if "\t" in line:
                rows: list[list[str]] = []
                while line_number <= len(lines) and "\t" in lines[line_number - 1]:
                    rows.append(lines[line_number - 1].split("\t"))
                    line_number += 1
                columns = max(len(row) for row in rows)
                table = document.add_table(rows=len(rows), cols=columns)
                table.style = "Table Grid"
                for table_row, values in zip(table.rows, rows):
                    for cell, value in zip(table_row.cells, values):
                        cell.text = value
            else:
                paragraph = document.add_paragraph()
                for text, tags in self._line_segments(line_number):
                    run = paragraph.add_run(text)
                    run.bold = "bold" in tags
                    run.italic = "italic" in tags
                    run.underline = "underline" in tags
                line_number += 1
        document.save(target)
        self.last_export = target
        self.status.configure(text=f"DOCX créé : {target.name}")

    def _export_pdf_to(self, target: Path) -> None:
        styles = getSampleStyleSheet()
        story = []
        table_rows: list[list[str]] = []
        lines = self._lines()
        for line_number, line in enumerate(lines + [""], start=1):
            if "\t" in line:
                table_rows.append(line.split("\t"))
                continue
            if table_rows:
                table = Table(table_rows, repeatRows=1)
                table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
                story.extend((table, Spacer(1, 4 * mm)))
                table_rows = []
            if line:
                fragments = []
                for text, tags in self._line_segments(line_number):
                    value = escape(text)
                    if "underline" in tags:
                        value = f"<u>{value}</u>"
                    if "italic" in tags:
                        value = f"<i>{value}</i>"
                    if "bold" in tags:
                        value = f"<b>{value}</b>"
                    fragments.append(value)
                content = "".join(fragments)
            else:
                content = "&nbsp;"
            story.append(Paragraph(content, styles["BodyText"]))
            story.append(Spacer(1, 2 * mm))
        SimpleDocTemplate(str(target), pagesize=A4, rightMargin=20 * mm, leftMargin=20 * mm, topMargin=20 * mm, bottomMargin=20 * mm).build(story)

    def _export_pdf(self) -> None:
        target = self._choose_export(".pdf", "Document PDF")
        if target is None:
            return
        self._export_pdf_to(target)
        self.last_export = target
        self.status.configure(text=f"PDF créé : {target.name}")

    def _print(self) -> None:
        temporary = Path(tempfile.gettempdir()) / "lou-codi-document-a-imprimer.pdf"
        self._export_pdf_to(temporary)
        open_path(temporary)
        messagebox.showinfo(
            "Impression",
            "Le PDF a été ouvert. Utilisez la commande Imprimer de votre lecteur PDF.",
            parent=self,
        )
