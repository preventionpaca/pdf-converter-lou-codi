from __future__ import annotations

import copy
import tempfile
from html import escape
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.text import WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt

from .canvas_document import (
    CanvasDocument,
    PageObject,
    export_canvas_pdf,
    flowed_text_context,
    pages_for_object,
    resolve_dynamic_text,
)

VML = "urn:schemas-microsoft-com:vml"
WORD = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
SEPARATE_DRAWING_KINDS = {
    "shape", "line", "polyline", "image", "mask", "dimension", "origin", "table",
}


def _text_style_at(item: PageObject, position: int) -> dict[str, object]:
    style: dict[str, object] = {
        "bold": item.bold,
        "italic": item.italic,
        "underline": item.underline,
        "strikethrough": item.strikethrough,
        "color": item.color,
        "highlight": "",
        "font_family": item.font_family,
        "font_size": item.font_size,
    }
    for run in item.text_styles:
        if int(run.get("start", 0)) <= position < int(run.get("end", 0)):
            style.update({key: value for key, value in run.items() if key not in {"start", "end"}})
    return style


def export_editable_docx(
    canvas_document: CanvasDocument,
    target: Path,
    *,
    compatibility_mode: bool = False,
) -> None:
    """Export either editable Word overlays or a universally visible page image."""
    background = copy.deepcopy(canvas_document)
    if not compatibility_mode:
        background.objects = [
            item for item in background.objects
            if item.kind != "text" and item.kind not in SEPARATE_DRAWING_KINDS
        ]
    with tempfile.TemporaryDirectory(prefix="lou-codi-docx-") as directory:
        temporary = Path(directory)
        pdf_path = temporary / "background.pdf"
        export_canvas_pdf(background, pdf_path)
        document = Document()
        section = document.sections[0]
        section.page_width = Inches(canvas_document.page_width / 72)
        section.page_height = Inches(canvas_document.page_height / 72)
        section.left_margin = section.right_margin = Inches(0)
        section.top_margin = section.bottom_margin = Inches(0)
        with pymupdf.open(pdf_path) as pdf:
            for page_number, page in enumerate(pdf):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                image_path = temporary / f"page-{page_number + 1}.png"
                page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False).save(image_path)
                paragraph.add_run().add_picture(str(image_path), width=Inches((canvas_document.page_width - 5) / 72))
                text_order = sorted(
                    enumerate(canvas_document.objects),
                    key=lambda indexed: (indexed[1].level, indexed[0]),
                )
                for index, item in text_order:
                    if (
                        not compatibility_mode
                        and item.kind == "text"
                        and page_number
                        in pages_for_object(item, canvas_document.page_count)
                    ):
                        document_name = Path(canvas_document.source_pdf).stem if canvas_document.source_pdf else target.stem
                        flowed_text, style_source, flow_offset = flowed_text_context(
                            item, canvas_document.objects
                        )
                        styled_item = copy.copy(item)
                        if style_source.object_id != item.object_id:
                            for attribute in (
                                "font_family", "font_size", "bold", "italic", "underline",
                                "strikethrough", "color", "align",
                            ):
                                setattr(styled_item, attribute, getattr(style_source, attribute))
                            flow_end = flow_offset + len(flowed_text)
                            styled_item.text_styles = [
                                {
                                    **run,
                                    "start": max(0, int(run.get("start", 0)) - flow_offset),
                                    "end": min(len(flowed_text), int(run.get("end", 0)) - flow_offset),
                                }
                                for run in style_source.text_styles
                                if int(run.get("start", 0)) < flow_end
                                and int(run.get("end", 0)) > flow_offset
                            ]
                            styled_item.text_links = [
                                {
                                    **run,
                                    "start": max(0, int(run.get("start", 0)) - flow_offset),
                                    "end": min(len(flowed_text), int(run.get("end", 0)) - flow_offset),
                                }
                                for run in style_source.text_links
                                if int(run.get("start", 0)) < flow_end
                                and int(run.get("end", 0)) > flow_offset
                            ]
                        link_relationships = [
                            (
                                int(link.get("start", 0)),
                                int(link.get("end", 0)),
                                document.part.relate_to(
                                    str(link.get("hyperlink", "")), RT.HYPERLINK, is_external=True
                                ),
                            )
                            for link in styled_item.text_links
                            if str(link.get("hyperlink", "")).strip()
                        ]
                        paragraph._p.append(
                            _textbox_xml(
                                styled_item,
                                page_number,
                                index,
                                canvas_document.page_count,
                                document_name,
                                flowed_text,
                                link_relationships,
                                None,
                            )
                        )
                    elif (
                        not compatibility_mode
                        and item.kind in SEPARATE_DRAWING_KINDS
                        and page_number in pages_for_object(item, canvas_document.page_count)
                    ):
                        object_image = temporary / f"object-{page_number + 1}-{index + 1}.png"
                        left, top, width, height = _render_object_png(item, object_image)
                        relationship_id, _image = document.part.get_or_add_image(str(object_image))
                        object_link_relationship = (
                            document.part.relate_to(item.hyperlink, RT.HYPERLINK, is_external=True)
                            if item.hyperlink.strip() else None
                        )
                        paragraph._p.append(
                            _picture_overlay_xml(
                                item, page_number, index, relationship_id,
                                left, top, width, height,
                                object_link_relationship,
                            )
                        )
                if page_number < pdf.page_count - 1:
                    paragraph.add_run().add_break(WD_BREAK.PAGE)
        target.parent.mkdir(parents=True, exist_ok=True)
        document.save(target)


def _render_object_png(item: PageObject, target: Path) -> tuple[float, float, float, float]:
    clone = copy.deepcopy(item)
    left = min(clone.x, clone.x + clone.width)
    top = min(clone.y, clone.y + clone.height)
    right = max(clone.x, clone.x + clone.width)
    bottom = max(clone.y, clone.y + clone.height)
    padding = 4.0
    width = max(8.0, right - left)
    height = max(8.0, bottom - top)
    clone.x -= left - padding
    clone.y -= top - padding
    clone.page = 0
    clone.pages = ""
    clone.hyperlink = ""
    miniature = CanvasDocument(
        page_width=width + padding * 2,
        page_height=height + padding * 2,
        objects=[clone],
    )
    pdf_path = target.with_suffix(".pdf")
    export_canvas_pdf(miniature, pdf_path)
    with pymupdf.open(pdf_path) as pdf:
        pixmap = pdf[0].get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=True)
        pixmap.save(target)
    return left - padding, top - padding, width + padding * 2, height + padding * 2


def _picture_overlay_xml(
    item: PageObject,
    page_number: int,
    index: int,
    relationship_id: str,
    left: float,
    top: float,
    width: float,
    height: float,
    hyperlink_relationship: str | None = None,
):
    drawing = "http://schemas.openxmlformats.org/drawingml/2006/main"
    picture = "http://schemas.openxmlformats.org/drawingml/2006/picture"
    word_drawing = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
    emu = 12700
    left_emu, top_emu = round(left * emu), round(top * emu)
    width_emu, height_emu = max(1, round(width * emu)), max(1, round(height * emu))
    drawing_id = page_number * 10000 + index + 1
    hyperlink = (
        f'<a:hlinkClick r:id="{hyperlink_relationship}"/>'
        if hyperlink_relationship else ""
    )
    xml = f"""
    <w:r xmlns:w="{WORD}"
         xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
         xmlns:wp="{word_drawing}" xmlns:a="{drawing}" xmlns:pic="{picture}">
      <w:drawing>
        <wp:anchor distT="0" distB="0" distL="0" distR="0"
          simplePos="0" relativeHeight="{max(0, item.level) + drawing_id}"
          behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">
          <wp:simplePos x="0" y="0"/>
          <wp:positionH relativeFrom="page"><wp:posOffset>{left_emu}</wp:posOffset></wp:positionH>
          <wp:positionV relativeFrom="page"><wp:posOffset>{top_emu}</wp:posOffset></wp:positionV>
          <wp:extent cx="{width_emu}" cy="{height_emu}"/>
          <wp:effectExtent l="0" t="0" r="0" b="0"/>
          <wp:wrapNone/>
          <wp:docPr id="{drawing_id}" name="LouCodiDrawing{page_number}_{index}">{hyperlink}</wp:docPr>
          <wp:cNvGraphicFramePr><a:graphicFrameLocks noChangeAspect="1"/></wp:cNvGraphicFramePr>
          <a:graphic>
            <a:graphicData uri="{picture}">
              <pic:pic>
                <pic:nvPicPr>
                  <pic:cNvPr id="{drawing_id}" name="LouCodiDrawing{page_number}_{index}">{hyperlink}</pic:cNvPr>
                  <pic:cNvPicPr/>
                </pic:nvPicPr>
                <pic:blipFill>
                  <a:blip r:embed="{relationship_id}"/>
                  <a:stretch><a:fillRect/></a:stretch>
                </pic:blipFill>
                <pic:spPr>
                  <a:xfrm><a:off x="0" y="0"/><a:ext cx="{width_emu}" cy="{height_emu}"/></a:xfrm>
                  <a:prstGeom prst="rect"><a:avLst/></a:prstGeom>
                </pic:spPr>
              </pic:pic>
            </a:graphicData>
          </a:graphic>
        </wp:anchor>
      </w:drawing>
    </w:r>
    """
    from docx.oxml import parse_xml

    return parse_xml(xml)


def _textbox_xml(
    item: PageObject,
    page_number: int,
    index: int,
    page_count: int,
    document_name: str,
    content: str,
    link_relationships: list[tuple[int, int, str]] | None = None,
    object_link_relationship: str | None = None,
):
    text = resolve_dynamic_text(content, page_number, page_count, document_name)
    alignment = {"left": "left", "center": "center", "right": "right", "justify": "both"}.get(item.align, "left")
    boundaries = {0, len(text)}
    for start, end, _relationship in link_relationships or []:
        boundaries.add(max(0, min(len(text), start)))
        boundaries.add(max(0, min(len(text), end)))
    for style_run in item.text_styles:
        boundaries.add(max(0, min(len(text), int(style_run.get("start", 0)))))
        boundaries.add(max(0, min(len(text), int(style_run.get("end", 0)))))
    fragments: list[str] = []
    ordered = sorted(boundaries)
    for start, end in zip(ordered, ordered[1:]):
        value = escape(text[start:end]).replace("\n", "</w:t><w:br/><w:t xml:space=\"preserve\">")
        if not value:
            continue
        relationship = next(
            (relationship for link_start, link_end, relationship in link_relationships or [] if link_start < end and link_end > start),
            None,
        )
        style = _text_style_at(item, start)
        family = {"Serif": "Liberation Serif", "Monospace": "Liberation Mono", "Cursive": "Comic Relief", "Manuscrite ancienne": "Z003", "Inter": "Inter", "JetBrains Mono": "JetBrains Mono", "Symboles": "DejaVu Sans", "Flèches": "DejaVu Sans"}.get(str(style["font_family"]), "DejaVu Sans")
        highlight = str(style.get("highlight", "")).lstrip("#")
        run_properties = (
            ("<w:b/>" if style["bold"] else "")
            + ("<w:i/>" if style["italic"] else "")
            + ("<w:u w:val=\"single\"/>" if style["underline"] else "")
            + ("<w:strike/>" if style["strikethrough"] else "")
            + f'<w:rFonts w:ascii="{family}" w:hAnsi="{family}"/>'
            + f'<w:sz w:val="{round(float(style["font_size"]) * 2)}"/>'
            + f'<w:color w:val="{str(style["color"]).lstrip("#")}"/>'
            + (f'<w:shd w:val="clear" w:color="auto" w:fill="{highlight}"/>' if highlight else "")
        )
        run = f'<w:r><w:rPr>{run_properties}{"<w:u w:val=\"single\"/><w:color w:val=\"1769AA\"/>" if relationship else ""}</w:rPr><w:t xml:space="preserve">{value}</w:t></w:r>'
        fragments.append(f'<w:hyperlink r:id="{relationship}">{run}</w:hyperlink>' if relationship else run)
    fragment_xml = "".join(fragments)
    fill = item.fill if item.fill.lower() not in {"", "none", "transparent"} else "#FFFFFF"
    filled = "f" if item.fill.lower() in {"", "none", "transparent"} else "t"
    stroked = "t" if item.border_width > 0 else "f"
    dash_style = "dot" if item.border_style == "dotted" else "dash" if item.border_style == "dashed" else "solid"
    run = f"""
    <w:r xmlns:w="{WORD}" xmlns:v="{VML}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <w:pict>
        <v:shape id="LouCodiTextBox{page_number}_{index}"
          style="position:absolute;margin-left:{item.x:.2f}pt;margin-top:{item.y:.2f}pt;width:{item.width:.2f}pt;height:{item.height:.2f}pt;rotation:{float(item.rotation) % 360:g};z-index:{item.level};mso-position-horizontal-relative:page;mso-position-vertical-relative:page"
          stroked="{stroked}" strokeweight="{max(0.0, item.border_width):g}pt"
          strokecolor="{item.border_color}" filled="{filled}" fillcolor="{fill}">
          <v:stroke dashstyle="{dash_style}"/>
          <v:textbox inset="3pt,2pt,3pt,2pt">
            <w:txbxContent>
              <w:p>
                <w:pPr><w:jc w:val="{alignment}"/></w:pPr>
                {fragment_xml}
              </w:p>
            </w:txbxContent>
          </v:textbox>
        </v:shape>
      </w:pict>
    </w:r>
    """
    xml = (
        f'<w:hyperlink xmlns:w="{WORD}" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="{object_link_relationship}">{run}</w:hyperlink>'
        if object_link_relationship else run
    )
    from docx.oxml import parse_xml

    return parse_xml(xml)
