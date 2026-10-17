from __future__ import annotations

from html import escape
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


LIGHTWEIGHT_SUFFIXES = {".docx", ".xlsx"}
EMU_PER_POINT = 12700


def convert_lightweight_office(source: Path, target: Path) -> None:
    """Render common modern Office documents without an office suite.

    This renderer deliberately targets ordinary technical documents. Complex floating
    objects, macros, SmartArt and exact printer metrics require a full office engine.
    """
    suffix = source.suffix.lower()
    target.parent.mkdir(parents=True, exist_ok=True)
    if suffix == ".docx":
        _docx_to_pdf(source, target)
    elif suffix == ".xlsx":
        _xlsx_to_pdf(source, target)
    else:
        raise ValueError(f"Le moteur léger ne prend pas en charge {source.suffix}.")


def _docx_to_pdf(source: Path, target: Path) -> None:
    document = Document(source)
    section = document.sections[0]
    page_size = (
        float(section.page_width) / EMU_PER_POINT,
        float(section.page_height) / EMU_PER_POINT,
    )
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "LouCodiBody",
        parent=styles["BodyText"],
        fontName="Helvetica",
        fontSize=10,
        leading=13,
        spaceAfter=4,
    )
    pdf = SimpleDocTemplate(
        str(target),
        pagesize=page_size,
        leftMargin=float(section.left_margin) / EMU_PER_POINT,
        rightMargin=float(section.right_margin) / EMU_PER_POINT,
        topMargin=float(section.top_margin) / EMU_PER_POINT,
        bottomMargin=float(section.bottom_margin) / EMU_PER_POINT,
        title=source.stem,
        author="Lou Còdi Workshop — moteur léger",
    )
    story = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = next((p for p in document.paragraphs if p._p is child), None)
            if paragraph is None:
                continue
            text = "".join(_run_markup(run) for run in paragraph.runs) or "&nbsp;"
            alignment = {
                1: TA_CENTER,
                2: TA_RIGHT,
                3: TA_JUSTIFY,
            }.get(paragraph.alignment, TA_LEFT)
            style = ParagraphStyle("p", parent=body, alignment=alignment)
            if paragraph.style and paragraph.style.name.startswith("Heading"):
                try:
                    level = int(paragraph.style.name.rsplit(" ", 1)[1])
                except (ValueError, IndexError):
                    level = 2
                style = ParagraphStyle(
                    "heading",
                    parent=styles["Heading1" if level == 1 else "Heading2"],
                    spaceBefore=7,
                    spaceAfter=5,
                )
            story.append(Paragraph(text, style))
            if child.xpath(".//w:br[@w:type='page']"):
                story.append(PageBreak())
        elif child.tag.endswith("}tbl"):
            table = next((t for t in document.tables if t._tbl is child), None)
            if table is not None:
                story.append(_docx_table(table, body))
                story.append(Spacer(1, 4 * mm))
    if not story:
        story.append(Paragraph("Document vide", body))
    pdf.build(story)


def _run_markup(run) -> str:
    text = escape(run.text).replace("\n", "<br/>")
    if not text:
        return ""
    if run.bold:
        text = f"<b>{text}</b>"
    if run.italic:
        text = f"<i>{text}</i>"
    if run.underline:
        text = f"<u>{text}</u>"
    return text


def _docx_table(table, body: ParagraphStyle) -> Table:
    rows = [
        [Paragraph(escape(cell.text).replace("\n", "<br/>"), body) for cell in row.cells]
        for row in table.rows
    ]
    result = Table(rows, repeatRows=1 if rows else 0, hAlign="LEFT")
    result.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#7D8FA3")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF5")),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return result


def _xlsx_to_pdf(source: Path, target: Path) -> None:
    workbook = load_workbook(source, data_only=True, read_only=False)
    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=8, leading=10)
    first_sheet = workbook.worksheets[0]
    page_size = landscape(A4) if first_sheet.page_setup.orientation == "landscape" else A4
    pdf = SimpleDocTemplate(
        str(target),
        pagesize=page_size,
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=source.stem,
        author="Lou Còdi Workshop — moteur léger",
    )
    story = []
    for sheet_index, sheet in enumerate(workbook.worksheets):
        if sheet_index:
            story.append(PageBreak())
        story.append(Paragraph(escape(sheet.title), styles["Heading2"]))
        rows = []
        for values in sheet.iter_rows(values_only=True):
            row = [Paragraph(escape("" if value is None else str(value)), cell_style) for value in values]
            if any(cell.getPlainText() for cell in row):
                rows.append(row)
        if not rows:
            story.append(Paragraph("Feuille vide", styles["BodyText"]))
            continue
        column_count = max(len(row) for row in rows)
        available = page_size[0] - 20 * mm
        widths = [available / column_count] * column_count
        table = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT", splitByRow=True)
        table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#8D99A6")),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DDE8F4")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 3),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        story.append(KeepTogether([table]) if len(rows) < 20 else table)
    pdf.build(story)
