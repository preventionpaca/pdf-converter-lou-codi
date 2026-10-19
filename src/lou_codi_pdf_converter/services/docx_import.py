from __future__ import annotations

import math
import re
import tempfile
import uuid
from collections import Counter, defaultdict, deque
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .canvas_document import (
    A4_HEIGHT,
    A4_WIDTH,
    CanvasDocument,
    PageObject,
    text_frame_line_capacity,
    text_frame_line_height,
    wrapped_text_lines,
)


def _font_family(name: str | None) -> str:
    lowered = (name or "").lower()
    if "courier" in lowered or "mono" in lowered:
        return "Monospace"
    if "times" in lowered or "serif" in lowered or "cambria" in lowered:
        return "Serif"
    return "Sans-serif"


_FONT_ATTRIBUTES = ("bold", "italic", "underline", "strike", "name", "size")


def _style_font_defaults(style: object) -> dict[str, object]:
    """Resolve a Word style chain once; large DOCX style tables make this costly."""
    values: dict[str, object] = {}
    current = style
    visited: set[str] = set()
    while current is not None:
        style_id = str(getattr(current, "style_id", "") or id(current))
        if style_id in visited:
            break
        visited.add(style_id)
        font = getattr(current, "font", None)
        for attribute in _FONT_ATTRIBUTES:
            value = getattr(font, attribute, None)
            if attribute not in values and value is not None:
                values[attribute] = value
        current = getattr(current, "base_style", None)
    return values


def _run_style(
    run: object,
    paragraph: object,
    *,
    paragraph_defaults: dict[str, object] | None = None,
    character_style_cache: dict[str, dict[str, object]] | None = None,
) -> dict[str, object]:
    font = getattr(run, "font", None)
    paragraph_values = paragraph_defaults
    if paragraph_values is None:
        paragraph_values = _style_font_defaults(getattr(paragraph, "style", None))
    character_values: dict[str, object] = {}
    style_id = str(getattr(getattr(run, "_r", None), "style", "") or "")
    if style_id:
        cache = character_style_cache if character_style_cache is not None else {}
        if style_id not in cache:
            cache[style_id] = _style_font_defaults(getattr(run, "style", None))
        character_values = cache[style_id]

    def inherited(attribute: str, fallback: object) -> object:
        direct = getattr(font, attribute, None)
        if direct is not None:
            return direct
        if attribute in character_values:
            return character_values[attribute]
        return paragraph_values.get(attribute, fallback)

    size = inherited("size", None)
    name = inherited("name", None)
    color = getattr(getattr(font, "color", None), "rgb", None)
    highlight = getattr(font, "highlight_color", None)
    highlight_names = {
        "YELLOW": "#FFF176",
        "BRIGHT_GREEN": "#A7F3A1",
        "TURQUOISE": "#8BE9FD",
        "PINK": "#F8BBD0",
        "BLUE": "#90CAF9",
        "RED": "#EF9A9A",
        "DARK_BLUE": "#5C6BC0",
        "TEAL": "#80CBC4",
        "GREEN": "#81C784",
        "VIOLET": "#D8B4FE",
        "DARK_RED": "#C62828",
        "DARK_YELLOW": "#F9A825",
        "GRAY_50": "#9E9E9E",
        "GRAY_25": "#D6D6D6",
        "BLACK": "#000000",
    }
    return {
        "bold": bool(inherited("bold", False)),
        "italic": bool(inherited("italic", False)),
        "underline": bool(inherited("underline", False)),
        "strikethrough": bool(inherited("strike", False)),
        "font_family": _font_family(str(name) if name else None),
        "font_size": float(size.pt) if size is not None else 12.0,
        "color": f"#{color}" if color is not None else "#000000",
        "highlight": highlight_names.get(getattr(highlight, "name", ""), ""),
    }


def _paragraph_prefix(paragraph: object, counters: dict[str, int]) -> str:
    properties = getattr(getattr(paragraph, "_p", None), "pPr", None)
    numbered = properties is not None and properties.numPr is not None
    style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "")
    if not numbered and "list" not in style_name.lower():
        return ""
    if "number" in style_name.lower():
        counters[style_name] = counters.get(style_name, 0) + 1
        return f"{counters[style_name]}. "
    return "• "


def _embedded_image_queues(source: object) -> dict[str, deque[dict[str, object]]]:
    """Extract inline or floating pictures, preserving their Word display size."""
    directory = Path(tempfile.mkdtemp(prefix="lou-codi-docx-import-"))
    images: dict[str, deque[dict[str, object]]] = defaultdict(deque)
    body = getattr(getattr(source, "element", None), "body", None)
    for blip in body.xpath(".//a:blip") if body is not None else ():
        try:
            relationship_id = blip.get(qn("r:embed"))
            part = source.part.related_parts[relationship_id]
            container = blip.getparent()
            while container is not None and container.tag not in {
                qn("wp:inline"), qn("wp:anchor")
            }:
                container = container.getparent()
            extents = container.xpath("./wp:extent") if container is not None else []
            extent = extents[0] if extents else None
            width = float(extent.get("cx")) / 12700.0 if extent is not None else 144.0
            height = float(extent.get("cy")) / 12700.0 if extent is not None else 96.0
            suffix = Path(str(getattr(part, "filename", ""))).suffix.lower()
            if not suffix:
                suffix = {
                    "image/jpeg": ".jpg",
                    "image/png": ".png",
                    "image/gif": ".gif",
                    "image/tiff": ".tiff",
                    "image/bmp": ".bmp",
                    "image/webp": ".webp",
                }.get(str(getattr(part, "content_type", "")).lower(), ".bin")
            target = directory / f"{uuid.uuid4().hex}{suffix}"
            target.write_bytes(part.blob)
            images[relationship_id].append(
                {
                    "path": str(target),
                    "width": max(1.0, width),
                    "height": max(1.0, height),
                }
            )
        except (AttributeError, KeyError, OSError, TypeError, ValueError):
            continue
    return images


def _table_cell_data(
    table: object,
    image_queues: dict[str, deque[dict[str, object]]],
) -> tuple[
    list[list[dict[str, object]]],
    list[list[int]],
    float,
    list[float],
    list[float],
]:
    """Convert a Word table to Lou Còdi's editable grid and merge metadata."""
    rows = list(getattr(table, "rows", ()))
    column_count = max((len(row.cells) for row in rows), default=1)
    cells: list[list[dict[str, object]]] = []
    # Key by the actual XML cell, not ``id(cell._tc)``.  The temporary
    # python-docx wrappers are released between rows and CPython can reuse an
    # integer id immediately, which used to merge unrelated cells and erase
    # their contents in complex forms.
    anchors: dict[object, tuple[int, int]] = {}
    coordinates: dict[object, set[tuple[int, int]]] = defaultdict(set)
    font_sizes: list[float] = []
    for row_index, row in enumerate(rows):
        output_row: list[dict[str, object]] = []
        for column_index in range(column_count):
            if column_index >= len(row.cells):
                output_row.append({"text": "", "hyperlink": "", "image_path": ""})
                continue
            cell = row.cells[column_index]
            cell_key = cell._tc
            coordinates[cell_key].add((row_index, column_index))
            anchor = anchors.setdefault(cell_key, (row_index, column_index))
            if anchor != (row_index, column_index):
                output_row.append({"text": "", "hyperlink": "", "image_path": ""})
                continue
            image_path = ""
            for blip in cell._tc.xpath(".//a:blip"):
                relationship_id = blip.get(qn("r:embed"))
                if relationship_id and image_queues[relationship_id]:
                    image_path = str(image_queues[relationship_id].popleft()["path"])
                    break
            horizontal = "left"
            for paragraph in cell.paragraphs:
                if paragraph.alignment is not None:
                    horizontal = {1: "center", 2: "right"}.get(
                        int(paragraph.alignment), "left"
                    )
                    break
            vertical = {
                WD_CELL_VERTICAL_ALIGNMENT.CENTER: "middle",
                WD_CELL_VERTICAL_ALIGNMENT.BOTTOM: "bottom",
            }.get(cell.vertical_alignment, "top")
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    size = getattr(run.font, "size", None)
                    if size is not None:
                        font_sizes.append(float(size.pt))
            visible_run = next(
                (
                    (paragraph, run)
                    for paragraph in cell.paragraphs
                    for run in paragraph.runs
                    if run.text.strip()
                ),
                None,
            )
            cell_style = (
                _run_style(
                    visible_run[1],
                    visible_run[0],
                    paragraph_defaults=_style_font_defaults(visible_run[0].style),
                )
                if visible_run
                else {}
            )
            cell_text = cell.text.lstrip("\n")
            cell_text_styles: list[dict[str, object]] = []
            search_at = 0
            for paragraph in cell.paragraphs:
                paragraph_defaults = _style_font_defaults(paragraph.style)
                for run in paragraph.runs:
                    if not run.text:
                        continue
                    start = cell_text.find(run.text, search_at)
                    if start < 0:
                        continue
                    end = start + len(run.text)
                    cell_text_styles.append({
                        "start": start,
                        "end": end,
                        **_run_style(
                            run,
                            paragraph,
                            paragraph_defaults=paragraph_defaults,
                        ),
                    })
                    search_at = end
            shading = cell._tc.tcPr.find(qn("w:shd")) if cell._tc.tcPr is not None else None
            shading_fill = shading.get(qn("w:fill"), "") if shading is not None else ""
            if shading_fill.lower() in {"", "auto", "nil"}:
                shading_fill = ""
            elif not shading_fill.startswith("#"):
                shading_fill = f"#{shading_fill}"
            if visible_run:
                direct_size = getattr(visible_run[1].font, "size", None)
                style_size = _style_font_defaults(visible_run[0].style).get("size")
                if direct_size is None and style_size is None:
                    # Word table templates often inherit an 8 pt base size.
                    # Treating an unspecified run as 12 pt made labels such as
                    # “N° SIRET” overflow across several rows.
                    cell_style["font_size"] = 8.0
            output_row.append(
                {
                    "text": cell_text,
                    "hyperlink": "",
                    "image_path": image_path,
                    "horizontal_align": horizontal,
                    "vertical_align": vertical,
                    "font_family": str(cell_style.get("font_family", "Sans-serif")),
                    "font_size": cell_style.get("font_size", 9.0),
                    "bold": bool(cell_style.get("bold", False)),
                    "italic": bool(cell_style.get("italic", False)),
                    "underline": bool(cell_style.get("underline", False)),
                    "color": str(cell_style.get("color", "#000000")),
                    "highlight": str(cell_style.get("highlight", "")),
                    "fill": shading_fill,
                    "text_styles": cell_text_styles,
                }
            )
        cells.append(output_row)
    merges: list[list[int]] = []
    for cell_key, positions in coordinates.items():
        if len(positions) <= 1:
            continue
        top = min(row for row, _column in positions)
        left = min(column for _row, column in positions)
        bottom = max(row for row, _column in positions)
        right = max(column for _row, column in positions)
        if positions == {
            (row, column)
            for row in range(top, bottom + 1)
            for column in range(left, right + 1)
        }:
            merges.append([top, left, bottom - top + 1, right - left + 1])
    typical_font_size = sorted(font_sizes)[len(font_sizes) // 2] if font_sizes else 9.0
    column_widths = [
        float(column.width.pt) if column.width is not None else 1.0
        for column in getattr(table, "columns", ())
    ]
    row_heights = [
        float(row.height.pt) if row.height is not None else 0.0
        for row in rows
    ]
    if not row_heights or any(height <= 0 for height in row_heights):
        row_heights = []
    return (
        cells or [[{"text": "", "hyperlink": "", "image_path": ""}]],
        merges,
        typical_font_size,
        column_widths,
        row_heights,
    )


def _estimated_table_height(cells: list[list[dict[str, str]]], font_size: float) -> float:
    """Estimate a readable equal-row table height for the canvas model."""
    row_heights: list[float] = []
    for row in cells:
        line_count = max(
            (max(1, cell.get("text", "").count("\n") + 1) for cell in row),
            default=1,
        )
        row_heights.append(max(24.0, line_count * font_size * 1.3 + 8.0))
    return max(40.0, sum(row_heights))


def _signature_date_values(text: str) -> list[str]:
    """Recognise repeated Word signature placeholders placed just below a table."""
    return [
        match.group(0).strip()
        for match in re.finditer(r"\bLe\s*[_ ]*[/]\s*[_ ]*[/]\s*20[_\d]*", text)
    ]


def _attach_signature_dates(table: PageObject, paragraph: Paragraph) -> bool:
    """Move a repeated date line into the bottom of the preceding signature cells."""
    if not table.table_cells:
        return False
    column_count = max(len(row) for row in table.table_cells)
    dates = _signature_date_values(paragraph.text)
    if len(dates) != column_count or column_count < 2:
        return False
    color = "#000000"
    for run in paragraph.runs:
        if run.text.strip():
            color = str(_run_style(run, paragraph).get("color", color))
            break
    bottom_row = len(table.table_cells) - 1
    for column, value in enumerate(dates):
        anchor_row, anchor_column = bottom_row, column
        for top, left, row_span, column_span in table.table_merges:
            if top <= bottom_row < top + row_span and left <= column < left + column_span:
                anchor_row, anchor_column = top, left
                break
        cell = table.table_cells[anchor_row][anchor_column]
        cell["bottom_text"] = value
        cell["bottom_color"] = color
    return True


def _imported_table_objects(
    source: object,
    image_queues: dict[str, deque[dict[str, object]]],
    *,
    page_count: int,
    margin_left: float,
    margin_top: float,
    width: float,
    height: float,
) -> list[PageObject]:
    """Place Word tables on their approximate source pages as editable objects."""
    tables = list(getattr(source, "tables", ()))
    if not tables:
        return []
    converted = [_table_cell_data(table, image_queues) for table in tables]
    assignments = [
        min(page_count - 1, index * page_count // len(tables))
        for index in range(len(tables))
    ]
    result: list[PageObject] = []
    for page in range(page_count):
        indices = [index for index, assigned in enumerate(assignments) if assigned == page]
        if not indices:
            continue
        estimates = [
            _estimated_table_height(converted[index][0], converted[index][2])
            for index in indices
        ]
        available = max(80.0, height - 12.0 * (len(indices) - 1))
        scale = min(1.0, available / max(1.0, sum(estimates)))
        cursor_y = margin_top
        for index, estimate in zip(indices, estimates, strict=True):
            cells, merges, font_size, column_widths, row_heights = converted[index]
            source_height = sum(row_heights) if row_heights else estimate
            table_height = max(40.0, source_height * scale)
            result.append(
                PageObject(
                    "table",
                    page,
                    margin_left,
                    cursor_y,
                    width,
                    table_height,
                    text="\n".join(
                        "\t".join(cell.get("text", "").replace("\n", " ") for cell in row)
                        for row in cells
                    ),
                    font_size=font_size,
                    table_cells=cells,
                    table_merges=merges,
                    table_column_widths=column_widths,
                    table_row_heights=row_heights,
                )
            )
            cursor_y += table_height + 12.0
    return result


def _dominant_text_metrics(
    styles: list[dict[str, object]],
) -> tuple[str, float]:
    """Choose a stable frame baseline from the character-weighted DOCX styles."""
    weights: Counter[tuple[str, float]] = Counter()
    for style in styles:
        try:
            length = max(0, int(style["end"]) - int(style["start"]))
            family = str(style.get("font_family", "Sans-serif"))
            size = float(style.get("font_size", 12.0))
        except (KeyError, TypeError, ValueError):
            continue
        weights[(family, size)] += length
    return weights.most_common(1)[0][0] if weights else ("Sans-serif", 12.0)


def _fit_nearly_empty_last_page(
    frame: PageObject, line_count: int
) -> tuple[int, int]:
    """Avoid a spurious page containing only a few metric-rounding lines."""
    capacity = max(1, text_frame_line_capacity(frame))
    page_count = max(1, math.ceil(line_count / capacity))
    if page_count <= 1:
        return capacity, page_count
    previous_pages = page_count - 1
    overflow = line_count - previous_pages * capacity
    tolerance = max(2, math.ceil(capacity * 0.05))
    if 0 < overflow <= tolerance:
        required_capacity = math.ceil(line_count / previous_pages)
        frame.font_size *= capacity / required_capacity
        capacity = max(1, text_frame_line_capacity(frame))
        page_count = max(1, math.ceil(line_count / capacity))
    return capacity, page_count


def import_docx_as_canvas(path: Path) -> CanvasDocument:
    """Import DOCX text, pictures, lists and formatting into editable objects."""
    source = Document(path)
    section = source.sections[0] if source.sections else None
    page_width = float(section.page_width.pt) if section and section.page_width else A4_WIDTH
    page_height = float(section.page_height.pt) if section and section.page_height else A4_HEIGHT
    margin_left = float(section.left_margin.pt) if section and section.left_margin else 42.0
    margin_right = float(section.right_margin.pt) if section and section.right_margin else 42.0
    margin_top = float(section.top_margin.pt) if section and section.top_margin else 42.0
    margin_bottom = float(section.bottom_margin.pt) if section and section.bottom_margin else 42.0
    width = max(120.0, page_width - margin_left - margin_right)
    height = max(100.0, page_height - margin_top - margin_bottom)

    parts: list[str] = []
    styles: list[dict[str, object]] = []
    alignments: list[dict[str, object]] = []
    image_placements: list[dict[str, object]] = []
    image_queues = _embedded_image_queues(source)
    counters: dict[str, int] = {}
    text_length = 0
    character_style_cache: dict[str, dict[str, object]] = {}
    paragraph_ranges: dict[object, tuple[int, int]] = {}
    alignment_names = {0: "left", 1: "center", 2: "right", 3: "justify", 4: "justify"}
    for paragraph in source.paragraphs:
        if parts:
            parts.append("\n")
            text_length += 1
        paragraph_start = text_length
        prefix = _paragraph_prefix(paragraph, counters)
        parts.append(prefix)
        text_length += len(prefix)
        paragraph_defaults = _style_font_defaults(getattr(paragraph, "style", None))
        for run in paragraph.runs:
            value = run.text.replace("\v", "\n").replace("\f", "\n")
            if value:
                start = text_length
                parts.append(value)
                end = start + len(value)
                text_length = end
                styles.append({
                    "start": start,
                    "end": end,
                    **_run_style(
                        run,
                        paragraph,
                        paragraph_defaults=paragraph_defaults,
                        character_style_cache=character_style_cache,
                    ),
                })
            for blip in run._r.xpath(".//a:blip"):
                relationship_id = blip.get(qn("r:embed"))
                if not relationship_id or not image_queues[relationship_id]:
                    continue
                image = image_queues[relationship_id].popleft()
                display_width = min(width, float(image["width"]))
                display_height = float(image["height"]) * display_width / float(image["width"])
                image_placements.append(
                    {
                        **image,
                        "width": display_width,
                        "height": display_height,
                        "position": text_length,
                    }
                )
                # Reserve the picture's vertical space in the linked text flow.
                spacer = "\n" * max(1, math.ceil(display_height / 15.0))
                parts.append(spacer)
                text_length += len(spacer)
        paragraph_end = text_length
        paragraph_ranges[paragraph._p] = (paragraph_start, paragraph_end)
        alignment = alignment_names.get(int(paragraph.alignment), "left") if paragraph.alignment is not None else "left"
        if alignment != "left" and paragraph_end > paragraph_start:
            alignments.append(
                {"start": paragraph_start, "end": paragraph_end, "align": alignment}
            )

    text = "".join(parts)
    base_family, base_size = _dominant_text_metrics(styles)
    prototype = PageObject(
        "text", 0, margin_left, margin_top, width, height,
        text=text, fill="transparent", border_width=0,
        font_family=base_family, font_size=base_size,
        text_styles=styles, paragraph_alignments=alignments,
    )
    # Plain documents keep a single linked text flow.  When tables are present,
    # however, a full-page flow followed by independently positioned tables
    # necessarily overlaps.  Compose the Word body in source order instead.
    if source.tables:
        converted_tables = {
            table._tbl: _table_cell_data(table, image_queues)
            for table in source.tables
        }
        blocks: list[tuple[str, object]] = []
        paragraph_group: list[Paragraph] = []

        def flush_paragraphs() -> None:
            if paragraph_group:
                blocks.append(("text", list(paragraph_group)))
                paragraph_group.clear()

        for block in source.iter_inner_content():
            if isinstance(block, Paragraph):
                paragraph_group.append(block)
                if len(_signature_date_values(block.text)) >= 2:
                    flush_paragraphs()
            elif isinstance(block, Table):
                flush_paragraphs()
                blocks.append(("table", block))
        flush_paragraphs()

        objects: list[PageObject] = []
        page = 0
        cursor_y = margin_top
        content_bottom = page_height - margin_bottom
        block_gap = 8.0

        def next_page() -> None:
            nonlocal page, cursor_y
            page += 1
            cursor_y = margin_top

        for kind, value in blocks:
            if kind == "table":
                (
                    cells,
                    merges,
                    table_font_size,
                    column_widths,
                    row_heights,
                ) = converted_tables[value._tbl]
                estimated = _estimated_table_height(cells, table_font_size)
                maximum_height = max(40.0, content_bottom - margin_top)
                table_height = min(
                    sum(row_heights) if row_heights else estimated,
                    maximum_height,
                )
                if cursor_y > margin_top and cursor_y + table_height > content_bottom:
                    next_page()
                objects.append(
                    PageObject(
                        "table",
                        page,
                        margin_left,
                        cursor_y,
                        width,
                        table_height,
                        text="\n".join(
                            "\t".join(
                                cell.get("text", "").replace("\n", " ")
                                for cell in row
                            )
                            for row in cells
                        ),
                        font_size=table_font_size,
                        table_cells=cells,
                        table_merges=merges,
                        table_column_widths=column_widths,
                        table_row_heights=row_heights,
                    )
                )
                cursor_y += table_height + block_gap
                if cursor_y >= content_bottom - 20.0:
                    next_page()
                continue

            paragraphs = list(value)
            first_visible = next(
                (index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip()),
                None,
            )
            if (
                objects
                and objects[-1].kind == "table"
                and first_visible is not None
                and _attach_signature_dates(objects[-1], paragraphs[first_visible])
            ):
                paragraphs.pop(first_visible)
                if not paragraphs:
                    continue
            ranges = [paragraph_ranges[paragraph._p] for paragraph in paragraphs]
            group_start, group_end = ranges[0][0], ranges[-1][1]
            group_text = text[group_start:group_end]
            if not group_text.strip():
                cursor_y = min(content_bottom, cursor_y + base_size * 1.35)
                continue

            group_styles = [
                {
                    **style,
                    "start": max(0, int(style["start"]) - group_start),
                    "end": min(group_end, int(style["end"])) - group_start,
                }
                for style in styles
                if int(style["end"]) > group_start and int(style["start"]) < group_end
            ]
            group_alignments = [
                {
                    **alignment,
                    "start": max(0, int(alignment["start"]) - group_start),
                    "end": min(group_end, int(alignment["end"])) - group_start,
                }
                for alignment in alignments
                if int(alignment["end"]) > group_start
                and int(alignment["start"]) < group_end
            ]
            group_family, group_size = _dominant_text_metrics(group_styles)
            exact_line_heights = [
                float(paragraph.paragraph_format.line_spacing.pt)
                for paragraph in paragraphs
                if getattr(paragraph.paragraph_format.line_spacing, "pt", None)
                is not None
            ]
            group_line_height = (
                sorted(exact_line_heights)[len(exact_line_heights) // 2]
                if exact_line_heights
                else 0.0
            )
            # TextMaker/Word compatibility documents commonly declare 8–8.5 pt
            # glyphs with an exact 8.5 pt leading.  Pillow/Tk measure those glyphs
            # wider than Word does, which otherwise creates an artificial extra
            # page.  Preserve the declared leading and compensate only the glyph
            # metrics used by Lou Còdi's editable text flow.
            if group_line_height and len(group_text) > 1000:
                metric_scale = 0.88
                group_size *= metric_scale
                for style in group_styles:
                    style["font_size"] = max(
                        1.0,
                        float(style.get("font_size", group_size)) * metric_scale,
                    )
            measuring_frame = PageObject(
                "text",
                page,
                margin_left,
                cursor_y,
                width,
                max(24.0, content_bottom - cursor_y),
                text=group_text,
                fill="transparent",
                border_width=0,
                font_family=group_family,
                font_size=group_size,
                text_styles=group_styles,
                paragraph_alignments=group_alignments,
                text_line_height=group_line_height,
            )
            line_count = max(1, len(wrapped_text_lines(group_text, measuring_frame)))
            line_height = text_frame_line_height(measuring_frame)
            remaining_lines = line_count
            group_frames: list[tuple[PageObject, int]] = []
            while remaining_lines > 0:
                available_height = content_bottom - cursor_y
                capacity = max(0, int((available_height - 8.0) / line_height))
                if capacity < 1:
                    next_page()
                    continue
                used_lines = min(remaining_lines, capacity)
                frame_height = min(
                    available_height,
                    max(24.0, used_lines * line_height + 8.0),
                )
                frame = PageObject(
                    "text",
                    page,
                    margin_left,
                    cursor_y,
                    width,
                    frame_height,
                    text=group_text if not group_frames else "",
                    fill="transparent",
                    border_width=0,
                    font_family=group_family,
                    font_size=group_size,
                    text_styles=group_styles if not group_frames else [],
                    paragraph_alignments=group_alignments if not group_frames else [],
                    text_line_height=group_line_height,
                )
                if group_frames:
                    group_frames[-1][0].flow_next_id = frame.object_id
                group_frames.append((frame, used_lines))
                objects.append(frame)
                remaining_lines -= used_lines
                cursor_y += frame_height + block_gap
                if remaining_lines > 0 or cursor_y >= content_bottom - 20.0:
                    next_page()

            group_images = [
                image
                for image in image_placements
                if group_start <= int(image["position"]) <= group_end
            ]
            for image in group_images:
                relative_position = max(0, int(image["position"]) - group_start)
                preceding_lines = max(
                    0,
                    len(wrapped_text_lines(group_text[:relative_position], measuring_frame)) - 1,
                )
                consumed = 0
                target_frame = group_frames[-1][0]
                line_on_frame = 0
                for frame, frame_lines in group_frames:
                    if preceding_lines < consumed + frame_lines:
                        target_frame = frame
                        line_on_frame = preceding_lines - consumed
                        break
                    consumed += frame_lines
                objects.append(
                    PageObject(
                        "image",
                        target_frame.page,
                        margin_left,
                        target_frame.y + line_on_frame * line_height,
                        float(image["width"]),
                        float(image["height"]),
                        image_path=str(image["path"]),
                    )
                )

        return CanvasDocument(
            page_width=page_width,
            page_height=page_height,
            page_count=max(1, max((item.page for item in objects), default=0) + 1),
            objects=objects,
        )

    line_count = max(1, len(wrapped_text_lines(text, prototype)))
    capacity, page_count = _fit_nearly_empty_last_page(prototype, line_count)
    frames = [prototype]
    for page in range(1, page_count):
        frames.append(PageObject(
            "text", page, margin_left, margin_top, width, height,
            fill="transparent", border_width=0,
        ))
        frames[-2].flow_next_id = frames[-1].object_id
    image_objects: list[PageObject] = []
    for image in image_placements:
        preceding_lines = max(
            0,
            len(wrapped_text_lines(text[: int(image["position"])], prototype)) - 1,
        )
        page = min(page_count - 1, preceding_lines // capacity)
        line_on_page = preceding_lines % capacity
        image_objects.append(
            PageObject(
                "image",
                page,
                margin_left,
                margin_top + line_on_page * 15.0,
                float(image["width"]),
                float(image["height"]),
                image_path=str(image["path"]),
            )
        )
    table_objects = _imported_table_objects(
        source,
        image_queues,
        page_count=page_count,
        margin_left=margin_left,
        margin_top=margin_top,
        width=width,
        height=height,
    )
    return CanvasDocument(
        page_width=page_width,
        page_height=page_height,
        page_count=page_count,
        objects=frames + image_objects + table_objects,
    )
