from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Callable


ProgressCallback = Callable[[str, int | None, int | None], None]


class _Pdf2DocxProgressHandler(logging.Handler):
    """Translate pdf2docx log records into stable, user-facing progress events."""

    _page_pattern = re.compile(r"\((\d+)/(\d+)\) Page (\d+)")

    def __init__(self, callback: ProgressCallback) -> None:
        super().__init__()
        self.callback = callback
        self.stage = "Préparation"

    def emit(self, record: logging.LogRecord) -> None:
        message = record.getMessage()
        if "[1/4]" in message:
            self.stage = "Ouverture du PDF"
            self.callback(self.stage, None, None)
        elif "[2/4]" in message:
            self.stage = "Analyse de la mise en page"
            self.callback(self.stage, None, None)
        elif "[3/4]" in message:
            self.stage = "Lecture des pages"
            self.callback(self.stage, None, None)
        elif "[4/4]" in message:
            self.stage = "Création du document Word"
            self.callback(self.stage, None, None)

        match = self._page_pattern.search(message)
        if match:
            current, total, _page_number = (int(value) for value in match.groups())
            self.callback(self.stage, current, total)


def convert_pdf_to_docx(
    source: Path,
    target: Path,
    password: str | None = None,
    progress: ProgressCallback | None = None,
) -> None:
    from pdf2docx import Converter

    converter = Converter(str(source), password or None)
    handler = _Pdf2DocxProgressHandler(progress) if progress else None
    logger = logging.getLogger()
    if handler:
        logger.addHandler(handler)
    try:
        # Multiprocessing can deadlock when pdf2docx runs from the packaged desktop
        # executable. Sequential conversion is predictable and still fast for the
        # documents handled by this application.
        converter.convert(str(target), multi_processing=False)
    finally:
        if handler:
            logger.removeHandler(handler)
        converter.close()


def convert_pdf_to_xlsx(source: Path, target: Path, password: str | None = None) -> int:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    count = 0
    with pdfplumber.open(str(source), password=password or None) as pdf:
        for page_number, page in enumerate(pdf.pages, 1):
            tables = page.extract_tables() or page.extract_tables(
                table_settings={
                    "vertical_strategy": "text",
                    "horizontal_strategy": "text",
                    "snap_tolerance": 3,
                    "intersection_tolerance": 5,
                }
            )
            for table in tables or []:
                if not table or not any(
                    any(str(cell or "").strip() for cell in row or []) for row in table
                ):
                    continue
                count += 1
                sheet = workbook.create_sheet(f"Page {page_number} - tableau {count}"[:31])
                widths: dict[int, int] = {}
                for row_index, row in enumerate(table, 1):
                    for column_index, cell in enumerate(row or [], 1):
                        value = str(cell or "").strip()
                        output_cell = sheet.cell(row_index, column_index, value)
                        output_cell.alignment = Alignment(wrap_text=True, vertical="top")
                        widths[column_index] = min(
                            55, max(widths.get(column_index, 10), len(value.split("\n")[0]) + 2)
                        )
                for column_index, width in widths.items():
                    sheet.column_dimensions[get_column_letter(column_index)].width = max(12, width)
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="244A70")
                sheet.freeze_panes = "A2"

    if not count:
        raise ValueError(
            "Aucun tableau exploitable détecté. Un PDF scanné nécessitera l’OCR local."
        )
    workbook.save(target)
    return count
